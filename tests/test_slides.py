import json

import pytest
from pptx import Presentation
from pptx.util import Inches

from slidegen.models import Figure, SlideContext, SlideSection
from slidegen.pipeline import rebuild_slides
from slidegen.slides import SLIDE_H, SLIDE_W, build_presentation, fit_font_size

LONG = "This is a fairly long bullet that describes a result with several technical terms and figures like 92.5% accuracy."
PICTURE = 13  # MSO_SHAPE_TYPE.PICTURE


def _ctx(make_image):
    wide = make_image("wide.png", (1200, 400))
    square = make_image("sq.png", (500, 450))
    return SlideContext("A Study of Things", "One-line tagline.", "paper.pdf", [
        SlideSection("intro", "Introduction", [LONG] * 3, "Notes for intro."),
        SlideSection("methods", "Methodology", [LONG] * 4, "Notes",
                     [Figure(wide, "Fig. 1: Workflow", "flowchart", "summary", "", 1200, 400)]),
        SlideSection("results", "Findings & Results", [LONG] * 5, "",
                     [Figure(square, "Fig. 2: Chart", "graph", "", "acc: RF 0.91 | XGB 0.93", 500, 450)]),
    ])


def _slide_text(slide):
    return " ".join(s.text_frame.text for s in slide.shapes if s.has_text_frame)


def test_deck_structure_16x9_title_agenda_content_closing(make_image, tmp_path):
    prs = Presentation(build_presentation(_ctx(make_image), tmp_path / "d.pptx"))
    assert prs.slide_width == Inches(SLIDE_W) and prs.slide_height == Inches(SLIDE_H)
    assert len(prs.slides) == 7          # title, agenda, intro, methods, wide-figure slide, results, closing
    texts = [_slide_text(s) for s in prs.slides]
    assert "A Study of Things" in texts[0] and "One-line tagline." in texts[0]
    assert "Introduction" in texts[1] and "Methodology" in texts[1]
    assert "Thank you" in texts[-1]


def test_every_shape_stays_inside_the_slide(make_image, tmp_path):
    prs = Presentation(build_presentation(_ctx(make_image), tmp_path / "d.pptx"))
    for i, slide in enumerate(prs.slides, 1):
        for shp in slide.shapes:
            assert shp.left >= 0 and shp.top >= 0, (i, shp.name)
            assert shp.left + shp.width <= prs.slide_width + 10, (i, shp.name)
            assert shp.top + shp.height <= prs.slide_height + 10, (i, shp.name)


def test_wide_figure_gets_own_slide_and_square_figure_sits_beside_text(make_image, tmp_path):
    prs = Presentation(build_presentation(_ctx(make_image), tmp_path / "d.pptx"))
    with_pic = [s for s in prs.slides if any(sh.shape_type == PICTURE for sh in s.shapes)]
    assert len(with_pic) == 2
    wide_slide, side_slide = with_pic
    assert LONG not in _slide_text(wide_slide) and "Fig. 1: Workflow" in _slide_text(wide_slide)
    assert LONG in _slide_text(side_slide) and "Fig. 2: Chart" in _slide_text(side_slide)
    assert "Chart values (approx.): acc: RF 0.91 | XGB 0.93" in _slide_text(side_slide)


def test_figures_keep_aspect_ratio(make_image, tmp_path):
    prs = Presentation(build_presentation(_ctx(make_image), tmp_path / "d.pptx"))
    pics = [sh for s in prs.slides for sh in s.shapes if sh.shape_type == PICTURE]
    assert pics
    for shp in pics:
        assert shp.width / shp.height == pytest.approx(shp.image.size[0] / shp.image.size[1], rel=0.01)


def test_speaker_notes_are_written(make_image, tmp_path):
    prs = Presentation(build_presentation(_ctx(make_image), tmp_path / "d.pptx"))
    assert prs.slides[2].notes_slide.notes_text_frame.text == "Notes for intro."


def test_bullets_are_real_bullets_not_prefix_characters(make_image, tmp_path):
    prs = Presentation(build_presentation(_ctx(make_image), tmp_path / "d.pptx"))
    xml = prs.slides[2].shapes._spTree.xml
    assert "a:buChar" in xml and "• " not in _slide_text(prs.slides[2])


def test_font_size_shrinks_with_more_text_and_none_when_hopeless():
    assert fit_font_size(["short bullet"] * 3, 11.9, 5.0) >= fit_font_size([LONG] * 5, 5.9, 5.0)
    assert fit_font_size([LONG * 6] * 8, 5.0, 3.0) is None


def test_overflowing_section_is_split_into_continuation_slides(tmp_path):
    ctx = SlideContext("T", "", "p.pdf", [SlideSection("r", "Results", [LONG * 3] * 5, "")])
    prs = Presentation(build_presentation(ctx, tmp_path / "d.pptx"))
    joined = [_slide_text(s) for s in prs.slides]
    assert any("Results (cont.)" in t for t in joined)


def test_rebuild_from_context_json_roundtrip(make_image, tmp_path):
    ctx = _ctx(make_image)
    path = tmp_path / "context.json"
    path.write_text(json.dumps(ctx.to_dict()), encoding="utf-8")
    assert SlideContext.from_dict(json.loads(path.read_text())) == ctx
    assert len(Presentation(rebuild_slides(path, tmp_path / "again.pptx")).slides) == 7


def test_figures_within_a_section_are_shown_in_figure_number_order(make_image, tmp_path):
    figs = [Figure(make_image(f"w{n}.png", (1200, 400)), f"Fig. {n}: Flow {n}", "flowchart", "", "", 1200, 400) for n in (2, 1)]
    ctx = SlideContext("T", "", "p.pdf", [SlideSection("methods", "Methodology", [LONG] * 3, "", figs)])
    prs = Presentation(build_presentation(ctx, tmp_path / "d.pptx"))
    texts = [_slide_text(s) for s in prs.slides]
    assert texts.index(next(t for t in texts if "Fig. 1" in t)) < texts.index(next(t for t in texts if "Fig. 2" in t))
